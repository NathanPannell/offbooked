"""Check a fixture's golden.json against its source file.

Usage:
  python -I verify.py <fixture_dir> [--sample] [--seed N] [--gaps N] [--quiet]

What it does (independently of the app's parser):
  1. Extracts raw text from source.pdf (PyMuPDF), source.docx (python-docx) or
     source.txt.
  2. Walks the golden speeches in order and aligns each one's word tokens to
     the source's word tokens inside a forward-moving window (difflib).
     Reports how many speeches are found almost verbatim, and lists the worst.
  3. Checks that one of the speaker's cue spellings appears just before each
     speech in the source.
  4. Lists long runs of source text between aligned speeches that the golden
     does not account for (should be stage directions, headers, notes; if it
     is dialogue, the golden is missing a speech).
  5. With --sample, prints the first 30, last 10 and 20 random speeches next
     to the raw source text they aligned to, plus characters and sections, for
     the manual spot check.
"""
from __future__ import annotations

import argparse
import difflib
import json
import os
import random
import re
import sys
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from benchlib import straight_quotes  # noqa: E402

TOKEN = re.compile(r"[a-z0-9]+")


def source_text(fixture: str) -> str:
    for name in os.listdir(fixture):
        if not name.startswith("source."):
            continue
        path = os.path.join(fixture, name)
        ext = name.rsplit(".", 1)[1]
        if ext == "pdf":
            import fitz  # PyMuPDF

            with fitz.open(path) as doc:
                return "\n".join(page.get_text(sort=True) for page in doc)
        if ext == "docx":
            import docx

            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append("\t".join(c.text for c in row.cells))
            return "\n".join(parts)
        if ext == "txt":
            raw = open(path, "rb").read()
            try:
                return raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                return raw.decode("latin-1")
    raise SystemExit(f"no source.* in {fixture}")


def fold(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return straight_quotes(text).lower().replace("'", "\u0000")


def tokens_with_offsets(text: str) -> list[tuple[str, int, int]]:
    # Apostrophes are removed for matching (don't -> dont) while keeping offsets.
    folded = fold(text)
    out = []
    for m in re.finditer(r"[a-z0-9\u0000]+", folded):
        tok = m.group().replace("\u0000", "")
        if tok:
            out.append((tok, m.start(), m.end()))
    return out


def toks(text: str) -> list[str]:
    return [t for t, _, _ in tokens_with_offsets(text)]


def match_in(g, src_toks, lo, hi, min_ratio):
    window = src_toks[lo:hi]
    if not window:
        return None
    sm = difflib.SequenceMatcher(None, g, window, autojunk=False)
    blocks = [b for b in sm.get_matching_blocks() if b.size]
    if not blocks:
        return None
    # Keep the densest cluster of matching blocks; stray matches far from the
    # main body (a common phrase a page later) are dropped.
    gap_limit = max(30, len(g))
    clusters, cur = [], [blocks[0]]
    for b in blocks[1:]:
        prev = cur[-1]
        if b.b - (prev.b + prev.size) > gap_limit:
            clusters.append(cur)
            cur = [b]
        else:
            cur.append(b)
    clusters.append(cur)
    blocks = max(clusters, key=lambda c: sum(b.size for b in c))
    ratio = sum(b.size for b in blocks) / len(g)
    if ratio < min_ratio:
        return None
    idxs = set()
    for b in blocks:
        idxs.update(range(lo + b.b, lo + b.b + b.size))
    return (ratio, lo + blocks[0].b, lo + blocks[-1].b + blocks[-1].size - 1, idxs)


def align(golden: dict, src_toks: list[str]):
    """Return per-speech (ratio, first_src_idx, last_src_idx, matched_src_idx_set).

    Pass 1 places speeches of 8+ words in order inside a forward window; pass 2
    places the short ones between their placed neighbours, so a common short
    line ("Well?") cannot drag the alignment forward.
    """
    n = len(src_toks)
    # Anchor: where the opening of the play (first ~40 golden tokens) sits in
    # the source, so front matter that quotes the first line is skipped.
    opening = [t for sp in golden["speeches"][:12] for t in toks(sp["text"])][:40]
    sm = difflib.SequenceMatcher(None, src_toks, opening, autojunk=False)
    best, start0 = 0, 0
    for start in range(0, max(1, n - len(opening)), max(1, len(opening) // 4)):
        sm.set_seq1(src_toks[start:start + 4 * len(opening)])
        score = sum(b.size for b in sm.get_matching_blocks())
        if score > best:
            best, start0 = score, start
    gtoks = [toks(sp["text"]) for sp in golden["speeches"]]
    results = [None] * len(gtoks)
    cursor = start0
    for i, g in enumerate(gtoks):
        if len(g) < 8:
            continue
        r = match_in(g, src_toks, cursor, min(n, cursor + 3 * len(g) + 1500), 0.6)
        if r:
            results[i] = r
            cursor = r[2] + 1
    lo = start0
    for i, g in enumerate(gtoks):
        if results[i] is not None:
            lo = results[i][2] + 1
            continue
        if not g:
            results[i] = (0.0, None, None, set())
            continue
        nxt = next((results[k][1] for k in range(i + 1, len(gtoks)) if results[k] is not None and results[k][1] is not None), n)
        hi = min(nxt, lo + 3 * len(g) + 600)
        r = match_in(g, src_toks, lo, max(hi, lo), 0.5)
        if r:
            results[i] = r
            lo = r[2] + 1
        else:
            results[i] = (0.0, None, None, set())
    return results


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("fixture")
    ap.add_argument("--sample", action="store_true")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--gaps", type=int, default=40, help="report unaccounted runs longer than N tokens")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    golden = json.load(open(os.path.join(args.fixture, "golden.json"), encoding="utf-8"))
    raw = source_text(args.fixture)
    src = tokens_with_offsets(raw)
    src_toks = [t for t, _, _ in src]
    results = align(golden, src_toks)

    speeches = golden["speeches"]
    cue_tokens: dict[str, list[list[str]]] = {}
    for c in golden["characters"]:
        forms = [c["name"]] + c["aliases"]
        cue_tokens[c["name"]] = [toks(f) for f in forms if toks(f)]

    found = sum(1 for r in results if r[1] is not None)
    exactish = sum(1 for r in results if r[0] >= 0.95)
    cue_ok = 0
    cue_miss = []
    for i, (ratio, first, _, _) in enumerate(results):
        if first is None:
            continue
        before = src_toks[max(0, first - 12):first]
        ok = any(
            any(before[j:j + len(ct)] == ct for j in range(len(before) - len(ct) + 1))
            for ct in cue_tokens[speeches[i]["speaker"]]
        )
        if ok:
            cue_ok += 1
        else:
            cue_miss.append(i)

    print(f"fixture: {args.fixture}")
    print(f"speeches: {len(speeches)}  aligned: {found}  >=95% verbatim: {exactish}  cue before speech: {cue_ok}/{found}")
    worst = sorted(range(len(results)), key=lambda i: results[i][0])[:15]
    if not args.quiet:
        print("\nlowest-matching speeches:")
        for i in worst:
            if results[i][0] >= 0.95:
                break
            print(f"  #{i} {results[i][0]:.2f} {speeches[i]['speaker']}: {speeches[i]['text'][:110]}")
        if cue_miss:
            print(f"\ncue not found before {len(cue_miss)} speeches, first 15:")
            for i in cue_miss[:15]:
                first = results[i][1]
                s = src[max(0, first - 12)][1]
                e = src[first][2]
                print(f"  #{i} {speeches[i]['speaker']}: ...{raw[s:e]!r}")

    # Unaccounted source runs between the first and last aligned speech.
    used = set()
    for r in results:
        used |= r[3]
    aligned = [r for r in results if r[1] is not None]
    if aligned:
        lo, hi = aligned[0][1], aligned[-1][2]
        run_start = None
        gaps = []
        for k in range(lo, hi + 1):
            if k not in used:
                if run_start is None:
                    run_start = k
            elif run_start is not None:
                if k - run_start > args.gaps:
                    gaps.append((run_start, k - 1))
                run_start = None
        print(f"\nunaccounted source runs > {args.gaps} tokens inside the play: {len(gaps)}")
        if not args.quiet:
            for a, b in gaps[:40]:
                snippet = re.sub(r"\s+", " ", raw[src[a][1]:src[b][2]])
                print(f"  [{b - a + 1} tok] {snippet[:240]}")
        head_unused = lo
        tail_unused = len(src_toks) - hi - 1
        print(f"source tokens before first speech: {head_unused}, after last speech: {tail_unused}")

    if args.sample:
        rnd = random.Random(args.seed)
        n = len(speeches)
        picks = list(range(min(30, n))) + list(range(max(0, n - 10), n))
        middle = [i for i in range(30, max(30, n - 10))]
        picks += sorted(rnd.sample(middle, min(20, len(middle))))
        print("\n=== spot check: golden vs source ===")
        for i in picks:
            ratio, first, last, _ = results[i]
            print(f"\n#{i} [{ratio:.2f}] {speeches[i]['speaker']}: {speeches[i]['text']}")
            if first is not None:
                s = src[max(0, first - 6)][1]
                e = src[last][2]
                print("   SRC: " + re.sub(r"\s+", " ", raw[s:e + 2])[:1200])
        print("\n=== characters ===")
        for c in golden["characters"]:
            count = sum(1 for s in speeches if s["speaker"] == c["name"])
            print(f"  {c['name']} {c['aliases']} ({count})")
        print("\n=== sections ===")
        print("  " + ", ".join(f"{s.get('act', '-')}.{s.get('scene', '-')}" for s in golden["sections"]))
        print(f"stageDirectionCount: {golden['stageDirectionCount']}")


if __name__ == "__main__":
    main()
